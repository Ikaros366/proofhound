#!/usr/bin/env python3
"""M6c 叙事事实守卫实战验收 demo（不进 pytest）。

对真实 engagement（缺省 eng-20260808T161941Z-ab7ca89a——M6c 背景案例：
T1 曾把 2 条 Reproduced 表述为"确认"）就地重建报告（narrative=true），
验证：
- 仓库自定义企业二进制模板 sha256 全程不变（本里程碑不碰；用 python-docx 打
  附录 B reason_cn 列的**临时副本**模拟建筑师同步后的模板做验收）；
- 初测综述中 F-2026-0004/0005（Reproduced）不再被表述为确认；确认计数
  表述 == 1；factguard 离线复验全语料零违规；
- 附录 B 出现中文归因（rejected_reasons_cn.json 非空且落进 docx）；
- 若发生修复重试，llm_repair_attempt 审计事件原文打印；
- default 模板 --no-llm 对照渲染回归（附录 B 同样显示中文归因）。

用法：
    .venv/bin/python scripts/demo_factguard.py [--engagement <dir>] [--env-file .env]
.env 需要：PROOFHOUND_T1_*（叙述生成）。
产物：engagement 目录内 report.docx / rejected_reasons_cn.json（就地重建，
与控制台"构建报告"等效）；模板副本落 evidence/demo_factguard/<时间戳>/（gitignored）。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))  # 允许直接以脚本方式运行

from docx import Document

from proofhound.findings import FindingState, FindingStore
from proofhound.report.__main__ import main as report_main
from proofhound.report.factguard import check_narrative_facts

ENTERPRISE_TEMPLATE_PATH = REPO_ROOT / "templates" / "custom_enterprise_template.docx"
DEFAULT_TEMPLATE = REPO_ROOT / "templates" / "default_template.docx"
DEFAULT_ENGAGEMENT = REPO_ROOT / "engagements" / "eng-20260808T161941Z-ab7ca89a"

CHECKS: list[tuple[str, bool]] = []


def check(name: str, ok: bool) -> None:
    CHECKS.append((name, ok))
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}")


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _make_synced_copy(src: Path, dst: Path) -> None:
    """复制自定义企业模板并把附录 B 排除原因列打上 reason_cn 回退（临时副本）。

    仓库二进制不动——本函数只写 dst（模拟建筑师同步后的模板形态）。
    docxtpl 标签可能被 Word 拆进多个 run，按段落整段重写（一次性副本，
    格式损失无所谓）。
    """
    doc = Document(str(src))
    replaced = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if "f.rejection_reason" in para.text:
                        new_text = para.text.replace(
                            "{{ f.rejection_reason }}",
                            "{{ f.reason_cn or f.rejection_reason }}",
                        )
                        for run in para.runs[1:]:
                            run.text = ""
                        para.runs[0].text = new_text
                        replaced += 1
    if replaced == 0:
        raise RuntimeError("自定义企业模板中未找到 {{ f.rejection_reason }} 单元格")
    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))
    print(f"[*] 模板同步副本: {dst}（改写 {replaced} 处，仓库模板不动）")


def _appendix_b_rows(doc: Document) -> list[list[str]]:
    table = next(
        t
        for t in doc.tables
        if t.rows[0].cells[0].text == "ID"
        and t.rows[0].cells[1].text == "漏洞类型"
        and t.rows[0].cells[3].text.strip() == "排除原因"
    )
    return [[c.text for c in r.cells] for r in table.rows[1:]]


def _sentences(text: str) -> list[str]:
    return [s for s in re.split(r"[。；;\n]+", text) if s.strip()]


def main() -> int:
    parser = argparse.ArgumentParser(description="ProofHound M6c 叙事事实守卫实战验收")
    parser.add_argument("--engagement", default=str(DEFAULT_ENGAGEMENT))
    parser.add_argument("--env-file", default=str(REPO_ROOT / ".env"))
    args = parser.parse_args()

    if not ENTERPRISE_TEMPLATE_PATH.exists():
        print("[skip] 企业模板缺失（公开仓库形态），跳过 factguard 演示："
              f"{ENTERPRISE_TEMPLATE_PATH}")
        return 0

    eng = Path(args.engagement)
    if not (eng / "findings.jsonl").is_file():
        print(f"[错误] engagement 无 findings.jsonl: {eng}", file=sys.stderr)
        return 2

    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    demo_dir = REPO_ROOT / "evidence" / "demo_factguard" / stamp
    demo_dir.mkdir(parents=True, exist_ok=True)

    # 0. 真实状态基线 + 仓库模板 sha256（验收末再核一遍）
    store = FindingStore(eng / "findings.jsonl")
    findings = store.load_all()
    states = {f.id: f.state for f in findings}
    confirmed = [f.id for f in findings if f.state is FindingState.CONFIRMED]
    reproduced = [f.id for f in findings if f.state is FindingState.REPRODUCED]
    rejected = [f.id for f in findings if f.state is FindingState.REJECTED]
    print(f"[*] 真实状态基线: confirmed={confirmed} reproduced={reproduced} "
          f"rejected={rejected}")
    template_sha_before = _sha256(ENTERPRISE_TEMPLATE_PATH)

    # 1. 自定义企业模板同步副本（仓库二进制不动）
    synced = demo_dir / "enterprise_synced.docx"
    _make_synced_copy(ENTERPRISE_TEMPLATE_PATH, synced)

    # 2. 就地重建报告（narrative=true，真实 T1；与控制台"构建报告"等效）
    print("\n[*] 就地重建报告（narrative=true，自定义企业同步副本模板）...")
    rc = report_main([
        "build", "--dir", str(eng), "--out", str(eng / "report.docx"),
        "--template", str(synced), "--env-file", args.env_file,
    ])
    if rc != 0:
        print(f"[失败] 报告构建退出码 {rc}（守卫双失败零写入也算构建失败）")
        return rc

    # 3. 修复重试审计事件原文打印（若发生）
    audit_lines = [
        json.loads(line)
        for line in (eng / "audit.jsonl").read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    repairs = [
        e for e in audit_lines
        if e.get("event") == "llm_repair_attempt" and e.get("caller") == "narrative"
    ]
    if repairs:
        print("\n[*] llm_repair_attempt 审计事件原文：")
        for event in repairs[-3:]:
            print("    " + json.dumps(event, ensure_ascii=False))
    else:
        print("\n[*] 本次构建未触发修复重试（首轮即过守卫）")

    # 4. 初测综述读回 + 离线复验
    sections = json.loads((eng / "narrative_sections.json").read_text(encoding="utf-8"))
    overview = sections.get("overview", "")
    remediation = sections.get("remediation", "")
    print("\n[*] 初测综述（overview）原文：")
    print("    " + overview.replace("\n", "\n    "))
    reasons_path = eng / "rejected_reasons_cn.json"
    reasons = json.loads(reasons_path.read_text(encoding="utf-8")) if reasons_path.is_file() else {}

    corpus = [overview, remediation, *reasons.values()]
    violations = check_narrative_facts(
        corpus,
        states,
        confirmed_count=len(confirmed),
        rejected_count=len(rejected),
    )
    check("factguard 离线复验全语料零违规", violations == [])
    if violations:
        for v in violations:
            print(f"    违规: {v}")

    fid_sentences = {
        fid: [s for s in _sentences(overview) if fid in s] for fid in reproduced
    }
    confirm_word = re.compile(r"(?<![未不无])确认|证实|confirm", re.IGNORECASE)
    check(
        f"F-2026-0004/0005（Reproduced）所在句无确认类词: {fid_sentences}",
        all(
            not confirm_word.search(s)
            for sentences in fid_sentences.values()
            for s in sentences
        ),
    )

    # 确认计数独立核对：凡与"确认"（非否定）共现于邻近窗口的 N 个/条/项，
    # 其 N 必须 == 真实 Confirmed 桶数（窗口双向：确认…N 单位 / N 单位…（为|被）确认）
    confirm_counts = []
    for match in re.finditer(r"([0-9]+|[一二三四五六七八九十])\s*[个条项]", overview):
        before = overview[max(0, match.start() - 8) : match.start()]
        after = overview[match.end() : match.end() + 6]
        if re.search(r"(?<![未不无])确认", before) or re.search(
            r"^(?:为|被)?[^。；，,]{0,3}?(?<![未不无])确认", after
        ):
            confirm_counts.append(match.group(1))
    check(
        f"确认计数表述 == 1（命中 {confirm_counts}）",
        bool(confirm_counts) and all(n in ("1", "一") for n in confirm_counts),
    )

    # 5. 附录 B 中文归因（自定义企业同步副本报告）
    print("\n[*] 误报中文归因（rejected_reasons_cn.json）原文：")
    for fid, text in sorted(reasons.items()):
        print(f"    {fid}: {text}")
    covered = set(reasons) & set(rejected)
    check(
        f"reason_cn 覆盖 rejected 桶（{len(covered)}/{len(rejected)}）",
        len(covered) == len(rejected) > 0,
    )

    doc = Document(str(eng / "report.docx"))
    rows = _appendix_b_rows(doc)
    print("\n[*] 附录 B 全行（自定义企业同步副本报告）：")
    for row in rows:
        print(f"    {row[0]} | {row[1]} | {row[3][:60]}")
    check(
        "附录 B 出现中文归因（行内容 == reasons_cn 文本）",
        any(row[3] in reasons.values() for row in rows),
    )

    # 6. default 模板 --no-llm 对照渲染回归
    print("\n[*] default 模板对照渲染（--no-llm 回归）...")
    default_out = demo_dir / "report_default.docx"
    rc = report_main([
        "build", "--dir", str(eng), "--out", str(default_out),
        "--template", str(DEFAULT_TEMPLATE), "--no-llm", "--env-file", args.env_file,
    ])
    if rc != 0:
        return rc
    default_rows = _appendix_b_rows(Document(str(default_out)))
    check(
        "default 模板附录 B 同样显示中文归因",
        any(row[3] in reasons.values() for row in default_rows),
    )

    # 7. 仓库自定义企业二进制模板全程未动
    check("仓库自定义企业模板 sha256 不变", _sha256(ENTERPRISE_TEMPLATE_PATH) == template_sha_before)

    failed = [name for name, ok in CHECKS if not ok]
    print(f"\n[*] 产物: {eng / 'report.docx'} / {default_out}")
    if failed:
        print(f"[失败] {len(failed)} 项未过: {failed}")
        return 1
    print("[*] 全部自检通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
