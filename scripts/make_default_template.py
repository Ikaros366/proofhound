#!/usr/bin/env python3
"""默认报告模板生成（M4，§5.7）：产出 templates/default_template.docx。

模板为 docxtpl（Jinja2 语法）全标签参考模板，兼作模板变量契约文档
（尾部附录列出全部可用变量）。固定章节骨架按 §5.7：测试概述 → 授权范围
→ 方法论 → 发现汇总表（按严重级）→ 详细发现（附证据索引与复现步骤）→
疑似未验证附录 → 已排除误报及原因附录 → 修复建议。

docxtpl 布局纪律（0.20.x 实测）：``{%tr ... %}`` 标签所在表格行会被整行
替换为标签——标签必须独占一行，数据行另起一行；``{%p ... %}`` 同理必须
独占段落；行内 ``{{ }}`` / ``{% if %}`` 不受限。

用法：
    .venv/bin/python scripts/make_default_template.py [--out <docx>]
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))  # 允许直接以脚本方式运行

DEFAULT_OUT = REPO_ROOT / "templates" / "default_template.docx"

#: 叙述槽位占位文案（--no-llm 或叙述未生成时渲染此文字）
PLACEHOLDER = "（本报告以 --no-llm 生成，叙述段落未生成）"

TABLE_STYLE = "Table Grid"


def _heading(doc: DocumentType, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _para(doc: DocumentType, text: str, style: str | None = None) -> None:
    doc.add_paragraph(text, style=style)


def _table(doc: DocumentType, headers: list[str]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = TABLE_STYLE
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    return table


def _loop_rows(table, for_tag: str, cells: list[str]) -> None:
    """三段式表格循环行：for 标签独占行 / 数据行 / endfor 标签独占行。"""
    ncols = len(table.columns)
    row_for = table.add_row()
    row_for.cells[0].text = for_tag
    row_data = table.add_row()
    for i, text in enumerate(cells):
        row_data.cells[i].text = text
    row_end = table.add_row()
    row_end.cells[0].text = "{%tr endfor %}"
    assert len(cells) == ncols, "数据行单元格数须与表头一致"


def _field_table(doc: DocumentType, pairs: list[tuple[str, str]]) -> None:
    """两列字段表（左标签右模板表达式），纯静态行。"""
    table = _table(doc, ["字段", "值"])
    for label, expr in pairs:
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = expr


def _evidence_index_table(doc: DocumentType) -> None:
    _para(doc, "证据包索引（证据原文落盘于证据包目录，可离线调出复核）：")
    _para(doc, "证据包目录：{{ f.evidence_pack.pack_dir }}")
    _para(doc, "{%p if not f.evidence_pack.assembled %}")
    _para(doc, "（证据包未组装：无 manifest.json）")
    _para(doc, "{%p endif %}")
    table = _table(doc, ["证据文件", "sha256", "来源锚点"])
    _loop_rows(
        table,
        "{%tr for e in f.evidence_pack.entries %}",
        [
            "{{ e.file or '（缺失）' }}",
            "{{ e.sha256 or '—' }}",
            "{{ e.source_ref }}",
        ],
    )


def build_document() -> DocumentType:
    """构建默认模板 Document（全标签参考模板）。"""
    doc = Document()

    _heading(doc, "渗透测试报告", level=0)
    _para(doc, "测试目标：{{ engagement.target or '（未记录）' }}")
    _para(doc, "授权范围：{{ engagement.scope or '（未记录）' }}")
    _para(
        doc,
        "测试时间窗：{{ engagement.started_at or '（未知）' }}"
        " ~ {{ engagement.finished_at or '（未知）' }}",
    )

    # ---- 1. 测试概述 ----
    _heading(doc, "1. 测试概述")
    _para(doc, "{{ sections.overview or '" + PLACEHOLDER + "' }}")

    # ---- 2. 授权范围 ----
    _heading(doc, "2. 授权范围")
    _para(
        doc,
        "本次测试在书面授权范围内进行。所有拟执行命令在沙箱中运行，"
        "目标逐条经 scope 白名单强制校验，越界拒绝并记审计日志。",
    )
    _para(doc, "授权范围（scope）：{{ engagement.scope or '（未记录）' }}")
    _para(doc, "测试目标：{{ engagement.target or '（未记录）' }}")
    _para(
        doc,
        "测试时间窗：{{ engagement.started_at or '（未知）' }}"
        " ~ {{ engagement.finished_at or '（未知）' }}",
    )

    # ---- 3. 方法论 ----
    _heading(doc, "3. 方法论")
    for line in [
        "候选发现（Signal）由扫描工具产出，经确定性 triage 建立假设（Hypothesis）；",
        "假设经 verify-* skill 行为验证（带会话 baseline 对照 + 工具行为确认）晋级 Reproduced；",
        "证据门按漏洞类型检查最低证据标准（行为证据为必要条件）；",
        "独立 Verifier 模型（与发现端不同模型）对抗校验通过后晋级 Confirmed；",
        "版本匹配型 CVE 与纯状态码证据永远禁止直接晋级 Confirmed（状态机铁律）；",
        "全部证据原文 100% 落盘证据包（含 sha256 清单），Confirmed 逐条附证据索引，可离线调出复核。",
    ]:
        _para(doc, line, style="List Bullet")

    # ---- 4. 发现汇总 ----
    _heading(doc, "4. 发现汇总")
    _para(
        doc,
        "Confirmed {{ summary.confirmed }} 条；需特定条件 {{ summary.conditional }} 条；"
        "疑似未验证 {{ summary.hypothesis }} 条；已排除误报 {{ summary.rejected }} 条。",
    )
    _para(doc, "Confirmed 按严重级分布：")
    table = _table(doc, ["严重级", "数量"])
    _loop_rows(
        table,
        "{%tr for sev, n in summary.severity_counts.items() %}",
        ["{{ sev }}", "{{ n }}"],
    )
    _para(doc, "Confirmed 发现汇总表（按严重级排序）：")
    _para(doc, "{%p if not confirmed_findings %}")
    _para(doc, "（无 Confirmed 发现）")
    _para(doc, "{%p endif %}")
    table = _table(doc, ["ID", "严重级", "漏洞类型", "资产", "标题"])
    _loop_rows(
        table,
        "{%tr for f in confirmed_findings %}",
        [
            "{{ f.id }}",
            "{{ f.severity }}",
            "{{ f.vuln_type }}",
            "{{ f.asset }}",
            "{{ f.title or '—' }}",
        ],
    )

    # ---- 5. 详细发现 ----
    _heading(doc, "5. 详细发现")
    _para(doc, "{%p if not confirmed_findings %}")
    _para(doc, "（无 Confirmed 发现）")
    _para(doc, "{%p endif %}")
    _para(doc, "{%p for f in confirmed_findings %}")
    _heading(doc, "5.{{ loop.index }} {{ f.title or f.vuln_type }}（{{ f.id }}）", level=2)
    _field_table(
        doc,
        [
            ("资产", "{{ f.asset }}"),
            ("参数", "{{ f.param or '—' }}"),
            ("严重级", "{{ f.severity }}"),
            ("漏洞类型", "{{ f.vuln_type }}"),
            (
                "前置条件",
                "{{ f.preconditions | join('；') if f.preconditions else '（无）' }}",
            ),
            ("置信度", "{{ f.confidence }}"),
        ],
    )
    _para(doc, "{%p if f.verification %}")
    _field_table(
        doc,
        [
            ("验证方法", "{{ f.verification.method }}"),
            ("验证者", "{{ f.verification.verified_by or '—' }}"),
            ("验证时间", "{{ f.verification.verified_at or '—' }}"),
        ],
    )
    _para(doc, "baseline 对照：{{ f.verification.baseline_diff or '—' }}")
    _para(doc, "{%p endif %}")
    _para(doc, "{%p if f.verifier %}")
    _para(
        doc,
        "Verifier 终审：{{ f.verifier.verdict }}（{{ f.verifier.model }}）"
        "——{{ f.verifier.reason }}",
    )
    _para(doc, "{%p endif %}")
    _para(doc, "风险叙述：")
    _para(doc, "{{ f.narrative or '" + PLACEHOLDER + "' }}")
    _para(doc, "{%p if f.verification %}")
    _para(doc, "复现步骤：")
    _para(doc, "{%p for step in f.verification.reproduction_steps %}")
    _para(doc, "{{ loop.index }}. {{ step }}")
    _para(doc, "{%p endfor %}")
    _para(doc, "{%p endif %}")
    _evidence_index_table(doc)
    _para(doc, "{%p endfor %}")

    # ---- 6. 需特定条件的发现 ----
    _heading(doc, "6. 需特定条件的发现（Reproduced，未晋级 Confirmed）")
    _para(
        doc,
        "以下发现已完成行为复现，但未通过全部确认环节（证据门或 Verifier 终审），"
        "按 fail-closed 原则停留 Reproduced，不作为 Confirmed 漏洞呈现。",
    )
    _para(doc, "{%p if not conditional_findings %}")
    _para(doc, "（无）")
    _para(doc, "{%p endif %}")
    _para(doc, "{%p for f in conditional_findings %}")
    _heading(doc, "6.{{ loop.index }} {{ f.title or f.vuln_type }}（{{ f.id }}）", level=2)
    _field_table(
        doc,
        [
            ("资产", "{{ f.asset }}"),
            ("参数", "{{ f.param or '—' }}"),
            ("严重级", "{{ f.severity }}"),
            ("漏洞类型", "{{ f.vuln_type }}"),
        ],
    )
    _para(doc, "{%p if f.verifier %}")
    _para(
        doc,
        "Verifier 终审：{{ f.verifier.verdict }}（{{ f.verifier.model }}）"
        "——{{ f.verifier.reason }}",
    )
    _para(doc, "{%p endif %}")
    _para(doc, "{{ f.narrative or '" + PLACEHOLDER + "' }}")
    _evidence_index_table(doc)
    _para(doc, "{%p endfor %}")

    # ---- 7. 疑似未验证（附录 A） ----
    _heading(doc, "7. 疑似未验证（附录 A）")
    _para(
        doc,
        "以下候选项未经行为验证，按“发现 ≠ 漏洞”原则仅供参考，"
        "不构成漏洞结论。",
    )
    _para(doc, "{%p if not hypothesis_findings %}")
    _para(doc, "（无）")
    _para(doc, "{%p endif %}")
    table = _table(doc, ["ID", "漏洞类型", "资产", "参数", "说明"])
    _loop_rows(
        table,
        "{%tr for f in hypothesis_findings %}",
        [
            "{{ f.id }}",
            "{{ f.vuln_type }}",
            "{{ f.asset }}",
            "{{ f.param or '—' }}",
            "{{ f.title or '—' }}",
        ],
    )

    # ---- 8. 已排除误报及原因（附录 B） ----
    _heading(doc, "8. 已排除误报及原因（附录 B）")
    _para(
        doc,
        "以下候选项经验证流程判定为误报或证据不足，逐条附排除原因，"
        "供复核与审计。",
    )
    _para(doc, "{%p if not rejected_findings %}")
    _para(doc, "（无）")
    _para(doc, "{%p endif %}")
    table = _table(doc, ["ID", "漏洞类型", "资产", "排除原因"])
    _loop_rows(
        table,
        "{%tr for f in rejected_findings %}",
        [
            "{{ f.id }}",
            "{{ f.vuln_type }}",
            "{{ f.asset }}",
            "{{ f.rejection_reason or '—' }}",
        ],
    )

    # ---- 9. 修复建议 ----
    _heading(doc, "9. 修复建议")
    _para(doc, "{{ sections.remediation or '" + PLACEHOLDER + "' }}")

    # ---- 附录 C：模板变量契约（文档用途） ----
    _heading(doc, "附录 C：模板变量契约（docxtpl/Jinja2）")
    _para(
        doc,
        "本模板同时是模板变量契约的参考实现。渲染环境为 Jinja2 "
        "StrictUndefined：引用未定义变量即报错。可用变量：",
    )
    for line in [
        "engagement.target / engagement.scope / engagement.started_at / engagement.finished_at",
        "summary.confirmed / summary.conditional / summary.hypothesis / summary.rejected / summary.severity_counts（dict，按严重级计数）",
        "confirmed_findings[] / conditional_findings[] / hypothesis_findings[] / rejected_findings[]，每项字段：",
        "  id, state, title, vuln_type, severity, asset, param, preconditions[], confidence, evidence_kinds[]",
        "  narrative（叙述槽位，可为 null）, rejection_reason（可为 null）",
        "  verification.{method, evidence_refs[], baseline_diff, reproduction_steps[], verified_by, verified_at}（可为 null）",
        "  verifier.{model, verdict, reason}（可为 null）",
        "  evidence_pack.{pack_dir, assembled, entries[]}；entries[] = {file, sha256, source_ref, line_anchor, missing}",
        "sections.overview / sections.remediation（固定章节叙述，可为 null）",
    ]:
        _para(doc, line, style="List Bullet")

    return doc


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="生成 ProofHound 默认报告模板")
    parser.add_argument(
        "--out",
        default=str(DEFAULT_OUT),
        help=f"输出 docx 路径（缺省 {DEFAULT_OUT}）",
    )
    args = parser.parse_args(argv)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    build_document().save(str(out))
    print(f"[+] 默认报告模板已生成: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
