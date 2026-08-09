#!/usr/bin/env python3
"""M4.5 模板适配验收 demo（不进 pytest）：自定义企业模板出真实报告。

链路：复制一份 demo_verify 产物（不改 M3b 原始产物）→ 补齐验收矩阵
（version-cve Hypothesis→Rejected；缺 web-exposure 则播种一条 Hypothesis；
engagement.json 补 extras：company_name/system_name/report_date）→
真实 T1 叙述版构建（--template 自定义企业模板）→ default_template 对照渲染
（回归不变）→ python-docx 读回自检：封面/时间/系统名、风险项循环与
Heading 4 连续、附录 A == 证据包 manifest、附录 B 含 version-cve、
1.4 流程章与源模板逐段一致。

用法：
    .venv/bin/python scripts/demo_report_enterprise.py [--dir evidence/demo_verify/<ts>]
                                                    [--dvwa-url http://127.0.0.1:8080]
.env 需要：PROOFHOUND_T1_*（叙述生成；缺则构建报错退出 2）。
产物落 evidence/demo_enterprise/<时间戳>/（gitignored）。
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))  # 允许直接以脚本方式运行
sys.path.insert(0, str(REPO_ROOT / "scripts"))  # 复用 seed_finding

from docx import Document

from proofhound.compliance.audit import AuditLog
from proofhound.findings import FindingState, FindingStore, assemble_evidence_pack
from proofhound.report.__main__ import main as report_main
from proofhound.report.data import build_context
from proofhound.report.render import _cn_date

import seed_finding

ENTERPRISE_TEMPLATE_PATH = REPO_ROOT / "templates" / "custom_enterprise_template.docx"
DEFAULT_TEMPLATE = REPO_ROOT / "templates" / "default_template.docx"

CHECKS: list[tuple[str, bool]] = []


def check(name: str, ok: bool) -> None:
    CHECKS.append((name, ok))
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}")


def _texts(doc: Document) -> tuple[list[str], list[str]]:
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return paras, cells


def _flow_section_texts(doc: Document) -> list[str]:
    """「渗透测试流程」章段落文本（该 Heading 2 → 下一 Heading 1 之间）。"""
    paras = doc.paragraphs
    start = next(
        i for i, p in enumerate(paras) if p.text.strip() == "渗透测试流程"
    )
    end = next(
        i
        for i, p in enumerate(paras[start + 1 :], start + 1)
        if p.style.name == "Heading 1"
    )
    return [p.text for p in paras[start:end]]


def prepare_matrix(evidence_dir: Path, dvwa_url: str) -> dict:
    """在副本上补齐验收矩阵 + engagement.json（含 extras），返回 extras。"""
    audit = AuditLog(evidence_dir / "audit.jsonl")
    store = FindingStore(evidence_dir / "findings.jsonl")

    version_cve = next(
        (f for f in store.load_all() if f.vuln_type == "version-cve"), None
    )
    if version_cve is not None and version_cve.state is FindingState.HYPOTHESIS:
        version_cve.audit = audit
        version_cve.transition(
            FindingState.REJECTED,
            actor="demo_report_enterprise",
            reason="版本匹配型 CVE：仅版本指纹比对，无行为验证手段，"
            "按 §5.4.1 铁律永远禁止直接 Confirmed，判误报排除",
        )
        store.append(version_cve)
        assemble_evidence_pack(version_cve, evidence_base=evidence_dir)
        print(f"[*] {version_cve.id} version-cve → Rejected（误报附录数据源）")

    if not any(f.vuln_type == "web-exposure" for f in store.load_all()):
        seed_finding.main([
            "--dir", str(evidence_dir), "--asset", f"{dvwa_url}/",
            "--vuln-type", "web-exposure", "--severity", "info",
            "--title", "DVWA 站点存活（web 暴露面）",
        ])

    extras = {
        "company_name": "演示客户单位",
        "system_name": "DVWA 演示系统",
        "report_date": "2026年8月",
    }
    engagement = {
        "target": dvwa_url,
        "scope": "127.0.0.0/8（仅 loopback，DVWA 本地靶场）",
        "started_at": "2026-08-07T02:00:00.000+00:00",
        "finished_at": "2026-08-07T04:30:00.000+00:00",
        **extras,
    }
    (evidence_dir / "engagement.json").write_text(
        json.dumps(engagement, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"[*] engagement.json 已写入（含 extras：{sorted(extras)}）")
    return engagement


def main() -> int:
    parser = argparse.ArgumentParser(description="ProofHound M4.5 自定义企业模板验收")
    parser.add_argument(
        "--dir",
        default=None,
        help="demo_verify 产物目录（缺省取 evidence/demo_verify/ 最新）",
    )
    parser.add_argument("--dvwa-url", default="http://127.0.0.1:8080")
    parser.add_argument("--env-file", default=str(REPO_ROOT / ".env"))
    args = parser.parse_args()

    if not ENTERPRISE_TEMPLATE_PATH.exists():
        print("[skip] 企业模板缺失（公开仓库形态），跳过自定义企业模板演示："
              f"{ENTERPRISE_TEMPLATE_PATH}")
        return 0

    if args.dir:
        src = Path(args.dir)
    else:
        candidates = sorted(
            p for p in (REPO_ROOT / "evidence" / "demo_verify").iterdir()
            if p.is_dir()
        )
        if not candidates:
            print("[错误] evidence/demo_verify/ 无产物，先跑 demo_verify_dvwa.py",
                  file=sys.stderr)
            return 2
        src = candidates[-1]
    if not (src / "findings.jsonl").is_file():
        print(f"[错误] 产物目录无 findings.jsonl: {src}", file=sys.stderr)
        return 2

    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    evidence_dir = REPO_ROOT / "evidence" / "demo_enterprise" / stamp
    evidence_dir.mkdir(parents=True, exist_ok=True)
    shutil.copytree(src, evidence_dir, dirs_exist_ok=True)
    print(f"[*] 验收副本: {evidence_dir}（源 {src}，原始产物未改动）")

    engagement = prepare_matrix(evidence_dir, args.dvwa_url)

    # 1. 真实 T1 叙述版（自定义企业模板；叙述槽位无守卫，必须叙述版构建）
    enterprise_out = evidence_dir / "report_enterprise.docx"
    print("\n[*] 构建自定义企业模板叙述版报告（T1 真实调用）...")
    rc = report_main([
        "build", "--dir", str(evidence_dir), "--out", str(enterprise_out),
        "--template", str(ENTERPRISE_TEMPLATE_PATH), "--env-file", args.env_file,
    ])
    if rc != 0:
        return rc

    # 2. default_template 对照（叙述已落盘，--no-llm 渲染回归）
    default_out = evidence_dir / "report_default.docx"
    print("\n[*] 构建 default_template 对照报告（渲染回归）...")
    rc = report_main([
        "build", "--dir", str(evidence_dir), "--out", str(default_out),
        "--template", str(DEFAULT_TEMPLATE), "--no-llm",
        "--env-file", args.env_file,
    ])
    if rc != 0:
        return rc

    # 3. 读回自检
    print("\n[*] 自定义企业报告读回自检 ...")
    store = FindingStore(evidence_dir / "findings.jsonl")
    confirmed = [
        f for f in store.load_all() if f.state is FindingState.CONFIRMED
    ]
    context = build_context(evidence_dir).as_template_context()

    doc = Document(str(enterprise_out))
    paras, cells = _texts(doc)

    # 3.1 封面/时间/系统名
    check(
        "封面/系统名/报告日期正确（extras 透传）",
        engagement["company_name"] in paras
        and engagement["system_name"] in paras
        and engagement["report_date"] in paras,
    )
    expected_time = (
        f"1）初测时间：{_cn_date(engagement['started_at'])}开始至"
        f"{_cn_date(engagement['finished_at'])}结束；"
    )
    check(f"初测时间渲染正确（cn_date：{expected_time}）", expected_time in paras)

    # 3.2 风险项循环展开 + Heading 4 连续（计数 + 桶序 + 【severity_cn】前缀）
    h4 = [p.text for p in doc.paragraphs if p.style.name == "Heading 4"]
    expected_h4 = [
        f"【{f['severity_cn']}】{f['title']}"
        for f in context["confirmed_findings"]
    ]
    check(
        f"风险项循环展开且 Heading 4 连续（{len(h4)} 条 == confirmed 数）",
        h4 == expected_h4,
    )

    # 3.2b M6b：CVSS 条件行——带分 Confirmed 逐条渲染（分数+向量原文），无分不出现
    scored = [
        f for f in context["confirmed_findings"] if f["cvss_score"] is not None
    ]
    expected_cvss = [
        f"CVSS：{f['cvss_score']}（{f['cvss_vector']}）" for f in scored
    ]
    cvss_lines = [p for p in paras if p.startswith("CVSS：")]
    check(
        f"CVSS 条件行逐条渲染（{len(scored)} 条带分 Confirmed）",
        cvss_lines == expected_cvss,
    )

    # 3.3 三段叙述槽位逐条在报告对应段
    parts_ok = all(
        f.narrative_parts
        and all(
            part in paras
            for part in (
                f.narrative_parts.description,
                f.narrative_parts.impact,
                f.narrative_parts.remediation,
            )
        )
        for f in confirmed
    )
    check("三段叙述（描述/危害/建议）逐条可回溯 finding", parts_ok)

    # 3.4 附录 A 与证据包 manifest 一致
    appendix_a = next(
        (t for t in doc.tables if t.rows[0].cells[0].text == "Finding ID"), None
    )
    index = context["evidence_index"]
    rows_ok = appendix_a is not None and len(appendix_a.rows) == len(index) + 1
    entries_ok = rows_ok and all(
        any(
            row[0] == e["finding_id"]
            and row[1] == (e["file"] or "")
            and row[2] == (e["sha256"] or "")
            and row[3] == e["source_ref"]  # source_ref 自带 #L 锚点，不重复拼接
            for row in ([c.text for c in r.cells] for r in appendix_a.rows[1:])
        )
        for e in index
    )
    check(
        f"附录 A 与证据包 manifest 一致（{len(index)} 条证据索引）",
        bool(entries_ok),
    )

    # 3.5 附录 B 含 version-cve
    rejected = next(f for f in store.load_all() if f.vuln_type == "version-cve")
    check(
        "附录 B 含 version-cve 及排除原因",
        any("version-cve" in cell for cell in cells)
        and any(rejected.rejection_reason[:20] in cell for cell in cells),
    )

    # 3.6 1.4 流程章与源模板逐段一致
    check(
        "1.4 渗透测试流程章与源模板逐段一致",
        _flow_section_texts(doc) == _flow_section_texts(Document(str(ENTERPRISE_TEMPLATE_PATH))),
    )

    # 3.7 default_template 回归：叙述派生段与证据索引照常渲染
    default_doc = Document(str(default_out))
    default_paras, default_cells = _texts(default_doc)
    narrative_ok = all(
        f.narrative and any(f.narrative in p for p in default_paras)
        for f in confirmed
    )
    summary_table = next(
        (t for t in default_doc.tables if t.rows[0].cells[0].text == "ID"
         and t.rows[0].cells[1].text == "严重级"),
        None,
    )
    check(
        "default_template 回归：派生叙述与汇总表照常渲染",
        narrative_ok
        and summary_table is not None
        and len(summary_table.rows) == len(confirmed) + 1,
    )

    failed = [name for name, ok in CHECKS if not ok]
    print(f"\n[*] 产物: {enterprise_out} / {default_out}")
    if failed:
        print(f"[失败] {len(failed)} 项未过: {failed}")
        return 1
    print("[*] 全部自检通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
