"""报告 CLI 测试（M4，§5.7）：python -m proofhound.report build。

--no-llm 端到端（tmp 出 docx、可打开、占位文字在）、配置/用法错误的
退出码、模块入口子进程冒烟。零真实网络。
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

from proofhound.report.__main__ import main as report_main

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "scripts"))  # 复用模板生成脚本

import make_default_template  # noqa: E402


@pytest.fixture
def template(tmp_path):
    path = tmp_path / "template.docx"
    make_default_template.main(["--out", str(path)])
    return path


def test_build_no_llm_end_to_end(report_evidence_dir, template, tmp_path):
    out = tmp_path / "report.docx"
    rc = report_main(
        [
            "build",
            "--dir",
            str(report_evidence_dir),
            "--out",
            str(out),
            "--template",
            str(template),
            "--no-llm",
        ]
    )
    assert rc == 0
    assert out.is_file()
    doc = Document(str(out))
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    assert any("渗透测试报告" in p for p in paras)
    assert any(make_default_template.PLACEHOLDER in p for p in paras)
    # 分桶计数进汇总表（2 条 confirmed + 表头）
    summary_table = next(
        t
        for t in doc.tables
        if t.rows[0].cells[0].text == "ID"
        and t.rows[0].cells[1].text == "严重级"
    )
    assert len(summary_table.rows) == 3


def test_build_missing_findings_store(tmp_path, template, capsys):
    rc = report_main(
        [
            "build",
            "--dir",
            str(tmp_path),
            "--out",
            str(tmp_path / "out.docx"),
            "--template",
            str(template),
            "--no-llm",
        ]
    )
    assert rc == 2
    assert "findings 存储不存在" in capsys.readouterr().err


def test_build_missing_template(report_evidence_dir, tmp_path, capsys):
    rc = report_main(
        [
            "build",
            "--dir",
            str(report_evidence_dir),
            "--out",
            str(tmp_path / "out.docx"),
            "--template",
            str(tmp_path / "nope.docx"),
            "--no-llm",
        ]
    )
    assert rc == 2
    assert "报告模板不存在" in capsys.readouterr().err


def test_module_entry_smoke(report_evidence_dir, template, tmp_path):
    """模块入口冒烟：python -m proofhound.report build --no-llm。"""
    out = tmp_path / "report.docx"
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "proofhound.report",
            "build",
            "--dir",
            str(report_evidence_dir),
            "--out",
            str(out),
            "--template",
            str(template),
            "--no-llm",
        ],
        capture_output=True,
        text=True,
        cwd=str(REPO_ROOT),
    )
    assert result.returncode == 0, result.stderr
    assert out.is_file()
    assert "报告已生成" in result.stdout
