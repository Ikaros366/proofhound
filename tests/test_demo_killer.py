"""M8d：demo_killer 纯函数罐头单测（无 Docker/LLM/浏览器）。

覆盖 demo_killer.py 抽出的可单测逻辑：L2 裁定策略、Confirmed 矩阵校验、
摘要表渲染、报告 XML 校验、脱敏自检。实弹链路不进 pytest（家族惯例）。
"""

from __future__ import annotations

import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "scripts"))  # 复用 demo 脚本（test_render.py 先例）

import demo_killer  # noqa: E402


def _canned_finding(
    fid: str,
    vuln_type: str,
    state: str = "confirmed",
    method: str | None = None,
    four_part: bool = True,
) -> dict:
    verification: dict = {"method": method or demo_killer.EXPECTED_METHOD[vuln_type]}
    if four_part:
        verification.update({"claim": "c", "expected": "e", "actual": "a"})
    return {
        "id": fid,
        "vuln_type": vuln_type,
        "state": state,
        "asset": f"http://127.0.0.1/{vuln_type}",
        "severity": "medium",
        "cvss_score": 5.0,
        "cvss_vector": "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:L/I:N/A:N",
        "verification": verification,
        "verifier": {"model": "m", "verdict": "confirm"},
    }


# ---- _decide：L2 确认裁定白名单 ----


def test_decide_approve_sqli_target_page():
    decision, _, mark = demo_killer._decide(
        "verify-sqli",
        "对 http://127.0.0.1:8080/vulnerabilities/sqli/?id=1 参数 id 执行 verify-sqli",
        frozenset(),
    )
    assert decision == "approve"
    assert mark == demo_killer.DVWA_SQLI_MARK


def test_decide_reject_sqli_other_page():
    decision, _, mark = demo_killer._decide(
        "verify-sqli",
        "对 http://127.0.0.1:8080/vulnerabilities/sqli_blind/?id=1 参数 id 执行",
        frozenset(),
    )
    assert decision == "reject"
    assert mark is None


def test_decide_approve_xss_and_idor_pair():
    assert demo_killer._decide(
        "verify-xss", "对 /vulnerabilities/xss_r/?name=1 参数 name 执行", frozenset()
    )[0] == "approve"
    assert demo_killer._decide(
        "verify-idor", "对 /invoice?id=1001 参数 id 执行", frozenset()
    )[0] == "approve"
    assert demo_killer._decide(
        "verify-idor", "对 /invoice?id=1002 参数 id 执行", frozenset()
    )[0] == "approve"


def test_decide_reject_idor_on_dvwa_and_duplicate_and_skip_l2():
    # DVWA 页面的 idor 候选（id= 但非 1001/1002）：拒绝
    assert demo_killer._decide(
        "verify-idor", "对 /vulnerabilities/sqli/?id=1 参数 id 执行", frozenset()
    )[0] == "reject"
    # 同一标记只批首条：第二条重复候选拒绝
    decision, _, mark = demo_killer._decide(
        "verify-idor", "对 /invoice?id=1001 参数 id 执行", frozenset({"id=1001"})
    )
    assert decision == "reject"
    assert mark is None
    # skip_l2 一律拒绝
    assert demo_killer._decide(
        "verify-xss", "对 /vulnerabilities/xss_r/?name=1 执行",
        frozenset(), skip_l2=True,
    )[0] == "reject"


# ---- _check_confirmed_matrix：需求①② ----


def test_check_confirmed_matrix_ok():
    findings = [
        _canned_finding("F-1", "sqli"),
        _canned_finding("F-2", "xss"),
        _canned_finding("F-3", "idor"),
        _canned_finding("F-4", "idor", state="rejected"),  # 对照组不计入
    ]
    assert demo_killer._check_confirmed_matrix(findings) == []


def test_check_confirmed_matrix_missing_type_and_count():
    findings = [_canned_finding("F-1", "sqli"), _canned_finding("F-2", "xss")]
    violations = demo_killer._check_confirmed_matrix(findings)
    assert any("Confirmed 数量" in v for v in violations)
    assert any("idor" in v for v in violations)


def test_check_confirmed_matrix_method_mismatch_and_four_part():
    bad = _canned_finding("F-9", "xss", method="sqlmap-confirmed", four_part=False)
    findings = [
        _canned_finding("F-1", "sqli"),
        bad,
        _canned_finding("F-3", "idor"),
    ]
    violations = demo_killer._check_confirmed_matrix(findings)
    assert any("不匹配" in v for v in violations)
    assert any("四段式缺 claim" in v for v in violations)
    weird = _canned_finding("F-8", "sqli", method="magic")
    violations = demo_killer._check_confirmed_matrix(findings + [weird])
    assert any("不在白名单" in v for v in violations)


# ---- _render_summary_table ----


def test_render_summary_table_contains_methods_assets_and_order():
    findings = [
        _canned_finding("F-3", "idor"),
        _canned_finding("F-1", "sqli"),
        _canned_finding("F-2", "xss"),
        _canned_finding("F-4", "idor", state="rejected"),  # 不进表
    ]
    files = {"F-1": ["sqlmap_stdout.txt"], "F-2": ["canary.json"], "F-3": ["j.json"]}
    table = demo_killer._render_summary_table(findings, files)
    assert "sqlmap-confirmed" in table
    assert "browser-confirmed" in table
    assert "dual-session-confirmed" in table
    assert "sqlmap_stdout.txt" in table and "canary.json" in table
    assert "F-4" not in table
    # 排序按 sqli → xss → idor
    assert table.index("F-1") < table.index("F-2") < table.index("F-3")


# ---- _check_report_xml：需求⑤ ----


def test_check_report_xml(tmp_path):
    import docx

    hit = tmp_path / "hit.docx"
    doc = docx.Document()
    doc.add_paragraph("验证方法：browser-confirmed；另有 dual-session-confirmed")
    doc.save(hit)
    assert demo_killer._check_report_xml(hit) == []

    miss = tmp_path / "miss.docx"
    doc2 = docx.Document()
    doc2.add_paragraph("只含 browser-confirmed")
    doc2.save(miss)
    assert demo_killer._check_report_xml(miss) == ["dual-session-confirmed"]

    # 文件不存在 = 全缺（fail-closed 方向）
    assert demo_killer._check_report_xml(tmp_path / "nope.docx") == [
        "browser-confirmed",
        "dual-session-confirmed",
    ]


# ---- _find_secret_leaks：需求⑥ ----


def test_find_secret_leaks_hit_and_miss(tmp_path):
    secret = "a1b2c3d4e5f60718"
    # 仅豁免文件含原文 → 通过
    (tmp_path / "session.json").write_text(f'{{"cookies": "phsess={secret}"}}')
    assert demo_killer._find_secret_leaks(tmp_path, [secret]) == []
    # 非豁免文件命中 → 报泄漏，且泄漏清单不回显原文（sha256 标记）
    (tmp_path / "audit.jsonl").write_text(f'{{"command": "curl -H {secret}"}}')
    leaks = demo_killer._find_secret_leaks(tmp_path, [secret])
    assert len(leaks) == 1
    assert "audit.jsonl" in leaks[0]
    assert secret not in leaks[0]
