"""模板渲染测试（M4，§5.7）：docxtpl + StrictUndefined + 默认模板读回。

覆盖：最小模板渲染、模板缺失/变量未定义的清晰报错、默认模板全量渲染
（章节齐全/汇总表行数/详细发现证据索引/误报附录/叙述占位）、渲染确定性
（同输入 → word/document.xml 内容一致）。python-docx 读回验证，零网络。
"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document

from proofhound.report.data import build_context
from proofhound.report.render import RenderError, render_docx

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "scripts"))  # 复用模板生成脚本

import make_default_template  # noqa: E402


@pytest.fixture
def minimal_template(tmp_path):
    """最小合法模板：一个变量 + 一个表格循环。"""
    doc = Document()
    doc.add_paragraph("目标：{{ engagement.target }}")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "ID"
    table.cell(0, 1).text = "类型"
    table.cell(1, 0).text = "{%tr for f in confirmed_findings %}"
    table.cell(2, 0).text = "{{ f.id }}"
    table.cell(2, 1).text = "{{ f.vuln_type }}"
    row = table.add_row()
    row.cells[0].text = "{%tr endfor %}"
    path = tmp_path / "minimal.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def default_template(tmp_path):
    """默认模板：由生成脚本现产（保证测试与脚本不漂移）。"""
    path = tmp_path / "default_template.docx"
    make_default_template.main(["--out", str(path)])
    return path


def _paragraph_texts(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs if p.text.strip()]


def _table_texts(doc: Document) -> list[str]:
    return [c.text for t in doc.tables for r in t.rows for c in r.cells]


# ---- 最小模板 ----


def test_minimal_template_renders(report_evidence_dir, minimal_template, tmp_path):
    context = build_context(report_evidence_dir).as_template_context()
    out = render_docx(context, minimal_template, tmp_path / "out.docx")
    doc = Document(str(out))
    assert any("目标：127.0.0.1:9" in p for p in _paragraph_texts(doc))
    # 汇总循环：表头 + 2 条 confirmed
    assert len(doc.tables[0].rows) == 3
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["F-2026-0005", "rce"]
    assert [c.text for c in doc.tables[0].rows[2].cells] == ["F-2026-0001", "sqli"]


def test_missing_template_raises(report_evidence_dir, tmp_path):
    context = build_context(report_evidence_dir).as_template_context()
    with pytest.raises(RenderError, match="模板不存在"):
        render_docx(context, tmp_path / "nope.docx", tmp_path / "out.docx")


def test_undefined_variable_raises(report_evidence_dir, minimal_template, tmp_path):
    """StrictUndefined：引用契约外变量即清晰报错，不静默空渲染。"""
    context = build_context(report_evidence_dir).as_template_context()
    del context["engagement"]  # 模板引用 engagement.target → 未定义
    with pytest.raises(RenderError, match="变量未定义"):
        render_docx(context, minimal_template, tmp_path / "out.docx")


# ---- 默认模板全量渲染 ----


def test_default_template_structure(
    report_evidence_dir, default_template, tmp_path
):
    context = build_context(report_evidence_dir).as_template_context()
    out = render_docx(context, default_template, tmp_path / "report.docx")
    doc = Document(str(out))
    paras = _paragraph_texts(doc)
    tables = _table_texts(doc)

    # 章节骨架齐全（§5.7 固定章节）
    for heading in (
        "1. 测试概述",
        "2. 授权范围",
        "3. 方法论",
        "4. 发现汇总",
        "5. 详细发现",
        "6. 需特定条件的发现",
        "7. 疑似未验证（附录 A）",
        "8. 已排除误报及原因（附录 B）",
        "9. 修复建议",
    ):
        assert any(heading in p for p in paras), f"缺章节: {heading}"

    # 汇总表：表头 + 2 条 confirmed（按 severity 排序：critical 在前）
    summary_table = next(
        t for t in doc.tables if t.rows[0].cells[0].text == "ID"
        and t.rows[0].cells[1].text == "严重级"
    )
    assert len(summary_table.rows) == 3
    assert summary_table.rows[1].cells[0].text == "F-2026-0005"
    assert summary_table.rows[2].cells[0].text == "F-2026-0001"

    # 详细发现：证据索引（文件名 + sha256）与复现步骤
    confirmed = next(
        f for f in context["confirmed_findings"] if f["id"] == "F-2026-0001"
    )
    entry = confirmed["evidence_pack"]["entries"][0]
    assert any(entry["sha256"] in cell for cell in tables)
    assert any(entry["file"] in cell for cell in tables)
    assert any("1. 步骤一" in p for p in paras)
    assert any("2. 步骤二" in p for p in paras)
    assert any("sqlmap-confirmed" in cell for cell in tables)

    # 误报附录：version-cve + 拒绝原因
    assert any("version-cve" in cell for cell in tables)
    assert any("铁律禁止直接 Confirmed" in cell for cell in tables)

    # 疑似未验证附录：web-exposure
    assert any("web-exposure" in cell for cell in tables)

    # 叙述槽位为空 → 占位文字（--no-llm 对照语义）
    assert any(make_default_template.PLACEHOLDER in p for p in paras)


def test_default_template_renders_narrative(
    report_evidence_dir, default_template, tmp_path
):
    """叙述落盘后渲染出叙述文字而非占位。"""
    import json

    from proofhound.findings.finding import FindingStore

    store = FindingStore(report_evidence_dir / "findings.jsonl")
    finding = store.get("F-2026-0001")
    finding.narrative = "该注入可致后台数据库内容泄漏。"
    store.append(finding)
    (report_evidence_dir / "narrative_sections.json").write_text(
        json.dumps(
            {"overview": "概述段落。", "remediation": "修复建议段落。"},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    context = build_context(report_evidence_dir).as_template_context()
    out = render_docx(context, default_template, tmp_path / "report.docx")
    paras = _paragraph_texts(Document(str(out)))
    assert any("该注入可致后台数据库内容泄漏。" in p for p in paras)
    assert any("概述段落。" == p for p in paras)
    assert any("修复建议段落。" == p for p in paras)


def test_render_deterministic(report_evidence_dir, default_template, tmp_path):
    """同输入同输出：两次渲染的 word/document.xml 内容字节一致。"""
    context = build_context(report_evidence_dir).as_template_context()
    out1 = render_docx(context, default_template, tmp_path / "r1.docx")
    out2 = render_docx(context, default_template, tmp_path / "r2.docx")

    def _document_xml(path: Path) -> bytes:
        with zipfile.ZipFile(path) as zf:
            return zf.read("word/document.xml")

    assert _document_xml(out1) == _document_xml(out2)


def test_render_escapes_xml_special_chars(report_evidence_dir, minimal_template, tmp_path):
    """替换值中的 & < > 必须原样保留（XML 转义，防 recover 解析静默吞字）。"""
    from proofhound.findings.finding import FindingStore

    store = FindingStore(report_evidence_dir / "findings.jsonl")
    finding = store.get("F-2026-0001")
    finding.asset = "http://127.0.0.1:9/app?id=1&Submit=Submit<b>"
    store.append(finding)
    context = build_context(report_evidence_dir).as_template_context()
    out = render_docx(context, minimal_template, tmp_path / "out.docx")
    cells = _table_texts(Document(str(out)))
    assert "sqli" in cells
    # 详细验证：默认模板下资产单元格原样呈现 & 与 <>
    default_tpl = tmp_path / "default.docx"
    make_default_template.main(["--out", str(default_tpl)])
    out2 = render_docx(context, default_tpl, tmp_path / "out2.docx")
    cells2 = _table_texts(Document(str(out2)))
    assert any(
        "http://127.0.0.1:9/app?id=1&Submit=Submit<b>" == cell for cell in cells2
    )


# ---- M4.5：cn_date 过滤器 + 自定义企业模板渲染读回 ----


@pytest.fixture
def cn_date_template(tmp_path):
    """cn_date 过滤器用最小模板。"""
    doc = Document()
    doc.add_paragraph("开始：{{ engagement.started_at | cn_date }}")
    doc.add_paragraph("结束：{{ engagement.finished_at | cn_date }}")
    path = tmp_path / "cn_date.docx"
    doc.save(str(path))
    return path


def test_cn_date_filter_iso(report_evidence_dir, cn_date_template, tmp_path):
    """ISO 时间 → 「2026年8月7日」（不补零；日期-only 串也可）。"""
    import json

    (report_evidence_dir / "engagement.json").write_text(
        json.dumps(
            {
                "started_at": "2026-08-07T12:34:56.789+00:00",
                "finished_at": "2026-08-09",
            }
        ),
        encoding="utf-8",
    )
    context = build_context(report_evidence_dir).as_template_context()
    out = render_docx(context, cn_date_template, tmp_path / "out.docx")
    paras = _paragraph_texts(Document(str(out)))
    assert "开始：2026年8月7日" in paras
    assert "结束：2026年8月9日" in paras


def test_cn_date_filter_empty_and_passthrough(tmp_path, cn_date_template):
    """空值 → 空串；非 ISO 串原样返回（用户手填的「2026年8月」类值可透）。"""
    import json

    (tmp_path / "findings.jsonl").write_text("", encoding="utf-8")
    (tmp_path / "engagement.json").write_text(
        json.dumps({"started_at": "", "finished_at": "2026年8月"}),
        encoding="utf-8",
    )
    context = build_context(tmp_path).as_template_context()
    out = render_docx(context, cn_date_template, tmp_path / "out.docx")
    paras = _paragraph_texts(Document(str(out)))
    assert "开始：" in paras  # 空值渲染为空串
    assert "结束：2026年8月" in paras  # 非 ISO 原样


ENTERPRISE_TEMPLATE_PATH = REPO_ROOT / "templates" / "custom_enterprise_template.docx"


def _flow_section_texts(doc: Document) -> list[str]:
    """「渗透测试流程」章段落文本（该 Heading 2 → 下一 Heading 1 之间）。"""
    paras = doc.paragraphs
    start = next(
        i for i, p in enumerate(paras) if p.text.strip() == "渗透测试流程"
    )
    end = next(
        i
        for i, p in enumerate(paras[start + 1 :], start + 1)
        if p.style.name == "Heading 1"
    )
    return [p.text for p in paras[start:end]]


def test_enterprise_template_renders(report_evidence_dir, tmp_path):
    """M4.5 验收：自定义企业模板渲染读回（封面/时间/风险项/附录 A B/流程章）。"""
    import json

    from proofhound.findings.finding import FindingStore, NarrativeParts

    store = FindingStore(report_evidence_dir / "findings.jsonl")
    parts_map = {
        "F-2026-0001": ("登录接口存在 SQL 注入。", "可致后台数据库内容泄漏。", "改用参数化查询。"),
        "F-2026-0005": ("接口可执行系统命令。", "可完全控制服务器。", "收敛危险函数并加白名单。"),
    }
    for fid, (desc, impact, remediation) in parts_map.items():
        finding = store.get(fid)
        finding.narrative_parts = NarrativeParts(
            description=desc, impact=impact, remediation=remediation
        )
        finding.narrative = f"{desc}\n{impact}\n{remediation}"
        store.append(finding)
    (report_evidence_dir / "narrative_sections.json").write_text(
        json.dumps({"overview": "概述段落。", "remediation": "建议段落。"}, ensure_ascii=False),
        encoding="utf-8",
    )
    (report_evidence_dir / "engagement.json").write_text(
        json.dumps(
            {
                "target": "http://127.0.0.1:9",
                "scope": "127.0.0.0/8",
                "started_at": "2026-08-07T01:00:00.000+00:00",
                "finished_at": "2026-08-09T02:00:00.000+00:00",
                "company_name": "某某单位",
                "system_name": "自定义企业演示系统",
                "report_date": "2026年8月",
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    context = build_context(report_evidence_dir).as_template_context()
    out = render_docx(context, ENTERPRISE_TEMPLATE_PATH, tmp_path / "enterprise.docx")
    doc = Document(str(out))
    paras = _paragraph_texts(doc)

    # 封面/时间/系统名（engagement extras 透传 + cn_date）
    assert "某某单位" in paras
    assert "自定义企业演示系统" in paras
    assert "2026年8月" in paras
    assert "1）初测时间：2026年8月7日开始至2026年8月9日结束；" in paras
    assert "概述段落。" in paras

    # 风险项循环：Heading 4 数量 == confirmed 数，severity_cn + 标题，桶序
    h4 = [p.text for p in doc.paragraphs if p.style.name == "Heading 4"]
    assert h4 == ["【严重】rce 标题", "【高】sqli 标题"]
    # 三段叙述槽位
    for text in (
        "登录接口存在 SQL 注入。",
        "可致后台数据库内容泄漏。",
        "改用参数化查询。",
        "接口可执行系统命令。",
    ):
        assert text in paras
    # {{r }} 富文本复现步骤：\n → <w:br/>（python-docx 读回为 \n）
    assert any("1. 步骤一\n2. 步骤二" in p for p in paras)

    # 附录 A：与证据包 manifest（evidence_index）逐条一致
    appendix_a = next(
        t for t in doc.tables if t.rows[0].cells[0].text == "Finding ID"
    )
    assert len(appendix_a.rows) == 1 + len(context["evidence_index"])
    row = [c.text for c in appendix_a.rows[1].cells]
    entry = context["evidence_index"][0]
    assert row[0] == entry["finding_id"] == "F-2026-0001"
    assert row[1] == entry["file"]
    assert row[2] == entry["sha256"]
    assert row[3] == entry["source_ref"]  # source_ref 自带 #L 锚点，模板不重复拼接
    assert entry["source_ref"].endswith(f"#L{entry['line_anchor']}")

    # 附录 B：version-cve + 排除原因
    appendix_b = next(
        t
        for t in doc.tables
        if t.rows[0].cells[0].text == "ID"
        and t.rows[0].cells[1].text == "漏洞类型"
    )
    assert len(appendix_b.rows) == 2
    cells = [c.text for c in appendix_b.rows[1].cells]
    assert cells[0] == "F-2026-0004" and cells[1] == "version-cve"
    assert "铁律禁止直接 Confirmed" in cells[3]

    # 1.4 渗透测试流程章：与源模板逐段一致（静态内容未被渲染改动）
    assert _flow_section_texts(doc) == _flow_section_texts(
        Document(str(ENTERPRISE_TEMPLATE_PATH))
    )


# ---- M6b：CVSS 行条件渲染 ----


def test_cvss_line_conditional_render(report_evidence_dir, default_template, tmp_path):
    """Confirmed 带分 → 渲染 CVSS 行；无分（旧 engagement 数据）→ 该行不出现。"""
    from proofhound.findings.finding import Finding, FindingStore

    store = FindingStore(report_evidence_dir / "findings.jsonl")
    store.append(
        Finding(
            id="F-2026-0010",
            state="confirmed",
            vuln_type="sqli",
            severity="critical",
            asset="http://127.0.0.1:9/app?q=10",
            dedup_key="sha256:cvss10",
            cvss_vector="CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H",
            cvss_score=9.8,
            created_at="2026-08-07T00:00:00.000+00:00",
            updated_at="2026-08-07T00:00:00.000+00:00",
        )
    )
    context = build_context(report_evidence_dir).as_template_context()
    out = render_docx(context, default_template, tmp_path / "out.docx")
    paras = _paragraph_texts(Document(str(out)))
    cvss_lines = [p for p in paras if p.startswith("CVSS：")]
    # 仅带分的 F-2026-0010 渲染；fixture 两条旧 confirmed（无分）不出现该行
    assert cvss_lines == [
        "CVSS：9.8（CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H）"
    ]
