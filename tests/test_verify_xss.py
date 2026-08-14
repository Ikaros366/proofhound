"""M8b verify-xss 测试：triage xss 分支、证据门 xss 白名单、_verify_xss
全链（FakeBrowser 预制 canary 事件，零真实浏览器）、多 verify skill 接线、
四段式证据结构与报告向后兼容。

覆盖验收点：hint 命中/零命中/独立上限/独立 triage_capped/审计键；
browser-confirmed 在 xss 白名单、不染指 sqli 白名单；FakeBrowser 预制
canary → 证据门 → Confirmed（四段式齐全、代码算分覆盖种子 severity）；
无 canary → REJECTED；浏览器错误 → blocked（fail-closed）；Verifier reject
→ REJECTED；runner 逐 skill 过闸；sqli 旧记录无四段式时报告渲染不炸。
"""

from __future__ import annotations

import json
import shutil
import sys
import time
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient

from proofhound.api import create_app
from proofhound.compliance.audit import AuditLog
from proofhound.compliance.scope import Scope
from proofhound.compliance.session import SessionConfig
from proofhound.core.orchestrator import Orchestrator
from proofhound.findings import (
    Finding,
    FindingState,
    FindingStore,
    compute_dedup_key,
)
from proofhound.llm.router import ModelRouter, Tier
from proofhound.report.data import build_context
from proofhound.report.render import render_docx
from proofhound.skills.registry import SkillRegistry
from proofhound.tools.sandbox import RunResult
from proofhound.verify.browser import PAYLOAD_TEMPLATES, BrowserProbeResult
from proofhound.verify.gate import check as gate_check

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "scripts"))  # 复用模板生成脚本（同 test_render 先例）

import make_default_template  # noqa: E402

BASE = "http://127.0.0.1:8080"
XSS_ASSET = f"{BASE}/vulnerabilities/xss_r/?name=1"
PHPSESSID = "df6a4b9c0e1f2a3b4c5d6e7f890abcde"
SESSION = SessionConfig(cookies={"PHPSESSID": PHPSESSID, "security": "low"})
# reflected XSS 典型向量（AV:N/AC:L/PR:N/UI:R/S:C/C:L/I:L/A:N → 6.1 medium）
CVSS_XSS = "CVSS:3.1/AV:N/AC:L/PR:N/UI:R/S:C/C:L/I:L/A:N"


# ---- triage：xss 分支 ----


class FakeLLM:
    def complete(self, messages):
        raise AssertionError("triage 不得调用 LLM")


def _param_row(asset, evidence_ref):
    return {
        "asset": asset,
        "status_code": 200,
        "title": None,
        "tech": [],
        "kind": "param-endpoint",
        "source_tool": "katana",
        "skill": "recon-crawl",
        "evidence_ref": evidence_ref,
    }


@pytest.fixture
def triage_env(tmp_path, make_skill_dir):
    evidence_dir = tmp_path / "evidence"
    evidence_dir.mkdir()
    raw = evidence_dir / "crawl.stdout.log"
    raw.write_text(
        "\n".join(f"line{i}" for i in range(1, 80)) + "\n", encoding="utf-8"
    )
    audit = AuditLog(evidence_dir / "audit.jsonl")
    orch = Orchestrator(
        SkillRegistry(make_skill_dir()).discover(),
        runner=SimpleNamespace(scope=Scope(networks=["127.0.0.0/8"])),
        llm=FakeLLM(),
        audit=audit,
        evidence_dir=evidence_dir,
    )
    return orch, audit, evidence_dir, raw


def _write_signals(evidence_dir, rows):
    signals_path = evidence_dir / "crawl.signals.jsonl"
    with signals_path.open("w", encoding="utf-8") as fh:
        for row in rows:
            fh.write(json.dumps(row) + "\n")


def _triage_summary(audit):
    return [e for e in audit.read_all() if e["event"] == "triage_completed"][0]


def test_xss_hint_creates_candidate(triage_env):
    """query 键命中 _XSS_PARAM_HINTS → 产 xss 候选；name 两表皆中 → sqli+xss 各一条。"""
    orch, audit, evidence_dir, raw = triage_env
    _write_signals(evidence_dir, [_param_row(f"{BASE}/xss_r/?name=1", f"{raw}#L6")])
    findings = orch.run_triage_phase()

    assert len(findings) == 2  # name 同产 sqli + xss（dedup 按 vuln_type 分量区分）
    by_type = {f.vuln_type: f for f in findings}
    xss = by_type["xss"]
    assert xss.param == "name"
    assert xss.severity == "medium"
    assert xss.evidence_kinds == ["crawl-endpoint"]
    assert by_type["sqli"].param == "name"
    assert by_type["sqli"].dedup_key != xss.dedup_key  # 互不合并
    summary = _triage_summary(audit)
    assert summary["created_by_type"] == {"sqli": 1, "xss": 1}
    assert summary["created_by_source"] == {"get_param": 2}


def test_xss_zero_hint_kept_signal(triage_env):
    """两表均不命中 → 保持 Signal。"""
    orch, audit, evidence_dir, raw = triage_env
    _write_signals(evidence_dir, [_param_row(f"{BASE}/x?foo=1", f"{raw}#L1")])
    assert orch.run_triage_phase() == []
    assert _triage_summary(audit)["kept_signal"] == 1


def test_xss_independent_cap_and_capped_audit(triage_env):
    """xss 独立上限 10（与 sqli 20 互不挤占）；独立 triage_capped 事件。"""
    orch, audit, evidence_dir, raw = triage_env
    rows = [_param_row(f"{BASE}/s{n}?id=1", f"{raw}#L1") for n in range(1, 26)]
    rows += [_param_row(f"{BASE}/x{n}?comment=1", f"{raw}#L2") for n in range(1, 13)]
    _write_signals(evidence_dir, rows)
    orch.run_triage_phase()

    store = FindingStore(evidence_dir / "findings.jsonl")
    findings = store.load_all()
    assert sum(1 for f in findings if f.vuln_type == "sqli") == 20  # sqli 满
    assert sum(1 for f in findings if f.vuln_type == "xss") == 10  # xss 独立上限
    capped = [e for e in audit.read_all() if e["event"] == "triage_capped"]
    by_type = {e["vuln_type"]: e for e in capped}
    assert by_type["sqli"]["dropped"] == 5
    assert by_type["xss"]["dropped"] == 2
    assert by_type["xss"]["limit"] == 10


def test_xss_out_of_scope_dropped(triage_env):
    """triage 层 scope 校验：xss 候选 asset 越界丢弃并记审计。"""
    orch, audit, evidence_dir, raw = triage_env
    _write_signals(
        evidence_dir, [_param_row("http://10.9.9.9:8080/x?comment=1", f"{raw}#L1")]
    )
    assert orch.run_triage_phase() == []
    oos = [e for e in audit.read_all() if e["event"] == "triage_out_of_scope"]
    assert len(oos) == 1


# ---- 证据门：xss 白名单（与 sqli 互不染指） ----


def _gate_finding(vuln_type, method):
    return Finding(
        id="F-2026-0001",
        state=FindingState.REPRODUCED,
        vuln_type=vuln_type,
        asset=XSS_ASSET,
        param="name",
        evidence_kinds=["crawl-endpoint", "behavioral"],
        verification={
            "method": method,
            "evidence_refs": ["evidence/x.log#L1"],
        },
        dedup_key="sha256:deadbeef",
        created_at="2026-08-14T00:00:00.000+00:00",
        updated_at="2026-08-14T00:00:00.000+00:00",
    )


def test_gate_xss_browser_confirmed_passes():
    assert gate_check(_gate_finding("xss", "browser-confirmed")).passed


def test_gate_xss_rejects_sqlmap_method():
    result = gate_check(_gate_finding("xss", "sqlmap-confirmed"))
    assert not result.passed
    assert any("白名单" in item for item in result.missing)


def test_gate_sqli_rejects_browser_method():
    """browser-confirmed 只进 xss 白名单，不染指 sqli。"""
    result = gate_check(_gate_finding("sqli", "browser-confirmed"))
    assert not result.passed


# ---- verify 全链：FakeBrowser + MockRouter ----


class FakeRunner:
    """按工具名弹出预制输出（httpx baseline）；写真实证据文件，零执行。"""

    def __init__(self, scope: Scope, evidence_dir: Path, script: dict[str, list]):
        self.scope = scope
        self.evidence_dir = evidence_dir
        self.script = {tool: list(responses) for tool, responses in script.items()}
        self.calls: list[tuple[str, list[str]]] = []
        self.egress_proxy_url = None
        self._seq = 0

    def run(self, tool, args, timeout=300, image=None):
        self.calls.append((tool, args))
        responses = self.script.get(tool) or []
        if not responses:
            raise AssertionError(f"FakeRunner 未预制 {tool} 的输出")
        stdout_text, exit_code = responses.pop(0)
        self._seq += 1
        stdout_path = self.evidence_dir / f"fake{self._seq}.stdout.log"
        stderr_path = self.evidence_dir / f"fake{self._seq}.stderr.log"
        stdout_path.write_text(stdout_text, encoding="utf-8")
        stderr_path.write_text("", encoding="utf-8")
        return RunResult(
            rejected=False,
            command=[tool, *args],
            exit_code=exit_code,
            stdout_path=stdout_path,
            stderr_path=stderr_path,
        )


class FakeBrowser:
    """预制 probe 结果的假浏览器（对齐 BrowserVerifier.probe 接口）。"""

    def __init__(self, evidence_dir: Path, *, canary_on=(), error_on=()):
        self.evidence_dir = evidence_dir
        self.canary_on = set(canary_on)
        self.error_on = set(error_on)
        self.calls: list[tuple] = []
        self.closed = False

    def probe(self, *, finding_id, seq, url, payload, token):
        self.calls.append((seq, url, payload, token))
        stem = f"xss_{finding_id}_{seq:02d}"
        paths = {}
        for key, suffix in (
            ("canary_path", "canary.json"),
            ("dom_path", "dom.html"),
            ("console_path", "console.json"),
            ("requests_path", "requests.json"),
        ):
            path = self.evidence_dir / f"{stem}_{suffix}"
            path.write_text("{}\n", encoding="utf-8")
            paths[key] = path
        if seq in self.error_on:
            return BrowserProbeResult(
                url=url, payload=payload, token=token,
                error="导航/驻留超时: fake timeout", **paths,
            )
        events = (
            [{"type": "alert", "token": token, "detail": token, "url": url,
              "payload": payload, "ts": 1}]
            if seq in self.canary_on
            else []
        )
        return BrowserProbeResult(
            url=url, payload=payload, token=token,
            events=events, canary=bool(events), **paths,
        )

    def close(self):
        self.closed = True


class MockRouter(ModelRouter):
    """罐头 T2 路由（继承 ModelRouter 过 ensure_router 的 isinstance 闸）。"""

    def __init__(self, reply: str):
        self.reply = reply  # 不调 super().__init__：不建 HTTP 客户端
        self.configs = {Tier.T2: SimpleNamespace(model="kimi-k3-test")}

    def complete(self, tier, messages):
        return self.reply


@pytest.fixture
def verify_env(tmp_path, make_skill_dir):
    evidence_dir = tmp_path / "evidence"
    evidence_dir.mkdir()
    audit = AuditLog(evidence_dir / "audit.jsonl")
    scope = Scope(networks=["127.0.0.0/8"], ports=[8080], session=SESSION)
    registry = SkillRegistry(
        make_skill_dir(name="verify-xss", tools=())
    ).discover()
    store = FindingStore(evidence_dir / "findings.jsonl")
    return SimpleNamespace(
        evidence_dir=evidence_dir, audit=audit, scope=scope,
        registry=registry, store=store,
    )


def _seed_xss(store: FindingStore, audit: AuditLog) -> Finding:
    finding = Finding(
        id=store.next_id(),
        state=FindingState.SIGNAL,
        vuln_type="xss",
        severity="medium",
        asset=XSS_ASSET,
        param="name",
        confidence="low",
        evidence_kinds=["crawl-endpoint"],
        dedup_key=compute_dedup_key(XSS_ASSET, "xss", "name"),
        source_signal_refs=[],
        created_at="2026-08-14T00:00:00.000+00:00",
        updated_at="2026-08-14T00:00:00.000+00:00",
        audit=audit,
    )
    finding.transition(FindingState.HYPOTHESIS, actor="triage", reason="测试种子")
    store.append(finding)
    return finding


def _httpx_ok_line():
    return json.dumps({"url": XSS_ASSET, "status_code": 200, "title": "DVWA"}) + "\n"


def _confirm_reply():
    return (
        '{"verdict": "confirm", "reason": "canary 执行事件链完整", '
        f'"cvss_vector": "{CVSS_XSS}", "cvss_rationale": "反射型 XSS 按证据定指标"}}'
    )


def _make_orch(env, browser, reply):
    runner = FakeRunner(env.scope, env.evidence_dir, {"httpx": [(_httpx_ok_line(), 0)]})
    orch = Orchestrator(
        env.registry, runner, MockRouter(reply), env.audit,
        evidence_dir=env.evidence_dir,
        browser_factory=lambda: browser,
    )
    return orch, runner


def test_verify_xss_full_chain_confirmed(verify_env):
    """canary 命中 → 证据门 → Verifier confirm → Confirmed（四段式齐全）。"""
    env = verify_env
    seed = _seed_xss(env.store, env.audit)
    browser = FakeBrowser(env.evidence_dir, canary_on={1})
    orch, runner = _make_orch(env, browser, _confirm_reply())

    processed = orch.run_verify_phase(skill_name="verify-xss")

    assert [f.id for f in processed] == [seed.id]
    assert [tool for tool, _ in runner.calls] == ["httpx"]  # baseline
    assert len(browser.calls) == 1  # 首发即命中即停
    seq, url, payload, token = browser.calls[0]
    assert token in payload and token in url  # canary token 内嵌
    assert payload in [t.replace("{token}", token) for t in PAYLOAD_TEMPLATES]

    finding = env.store.load_all()[0]
    assert finding.state is FindingState.CONFIRMED
    verification = finding.verification
    assert verification.method == "browser-confirmed"
    assert "behavioral" in finding.evidence_kinds
    # 四段式
    assert verification.claim == "参数 name 的输入在浏览器中被执行"
    assert "canary token" in verification.expected
    assert token in verification.actual
    # M6b 机制：代码算分覆盖种子 severity
    assert finding.cvss_vector == CVSS_XSS
    assert finding.cvss_score == 6.1
    assert finding.severity == "medium"
    # 证据引用齐备（baseline + canary/dom/console/requests）
    assert len(verification.evidence_refs) == 5
    assert any("canary.json" in ref for ref in verification.evidence_refs)
    # 审计：按次计数 + 门过 + Verifier + verify_completed
    events = env.audit.read_all()
    attempts = [e for e in events if e["event"] == "xss_probe_attempt"]
    assert len(attempts) == 1 and attempts[0]["canary"] is True
    assert attempts[0]["token"] == token
    verdicts = [e for e in events if e["event"] == "verifier_verdict"]
    assert verdicts and verdicts[0]["cvss_vector"] == CVSS_XSS
    completed = [e for e in events if e["event"] == "verify_completed"]
    assert completed and completed[0]["confirmed"] == 1
    assert browser.closed  # phase 收尾释放浏览器
    # 证据包含 canary/dom 条目
    manifest = json.loads(
        (env.evidence_dir / "findings" / seed.id / "manifest.json").read_text(
            encoding="utf-8"
        )
    )
    names = [item["file"] for item in manifest["items"]]
    assert any(n and "canary" in n for n in names)  # 包内文件名带 sha256 中缀
    assert any(n and "dom" in n for n in names)


def test_verify_xss_no_canary_rejected(verify_env):
    """对照组：全部 payload 干净完成且无 canary → REJECTED，不得 Confirmed。"""
    env = verify_env
    _seed_xss(env.store, env.audit)
    browser = FakeBrowser(env.evidence_dir)  # 无 canary
    orch, _runner = _make_orch(env, browser, "{}")

    orch.run_verify_phase(skill_name="verify-xss")

    finding = env.store.load_all()[0]
    assert finding.state is FindingState.REJECTED
    assert "canary" in finding.rejection_reason
    assert len(browser.calls) == len(PAYLOAD_TEMPLATES)  # 全量尝试
    attempts = [
        e for e in env.audit.read_all() if e["event"] == "xss_probe_attempt"
    ]
    assert len(attempts) == len(PAYLOAD_TEMPLATES)  # 按次计数审计
    assert all(e["canary"] is False for e in attempts)


def test_verify_xss_browser_error_blocked_fail_closed(verify_env):
    """浏览器错误且未命中 canary → blocked（覆盖不全不驳回），停留 Hypothesis。"""
    env = verify_env
    _seed_xss(env.store, env.audit)
    browser = FakeBrowser(env.evidence_dir, error_on={2})  # 第 2 次出错
    orch, _runner = _make_orch(env, browser, "{}")

    orch.run_verify_phase(skill_name="verify-xss")

    finding = env.store.load_all()[0]
    assert finding.state is FindingState.HYPOTHESIS  # fail-closed 停留原态
    blocked = [e for e in env.audit.read_all() if e["event"] == "verify_blocked"]
    assert blocked and "覆盖不全" in blocked[-1]["reason"]


def test_verify_xss_verifier_reject(verify_env):
    """canary 命中但 Verifier reject → REJECTED(actor=verifier)。"""
    env = verify_env
    _seed_xss(env.store, env.audit)
    browser = FakeBrowser(env.evidence_dir, canary_on={1})
    orch, _runner = _make_orch(
        env, browser, '{"verdict": "reject", "reason": "证据链不足以支撑"}'
    )

    orch.run_verify_phase(skill_name="verify-xss")

    finding = env.store.load_all()[0]
    assert finding.state is FindingState.REJECTED
    assert finding.verifier.verdict == "reject"


# ---- 报告层：四段式透传 + 向后兼容 ----


def test_four_part_fields_in_report_context(verify_env, tmp_path):
    """xss Confirmed 的 context.verification 四段式齐全；默认模板渲染出现四段式段落。"""
    env = verify_env
    _seed_xss(env.store, env.audit)
    browser = FakeBrowser(env.evidence_dir, canary_on={1})
    orch, _runner = _make_orch(env, browser, _confirm_reply())
    orch.run_verify_phase(skill_name="verify-xss")

    context = build_context(env.evidence_dir)
    confirmed = context.confirmed_findings[0]
    verification = confirmed.verification
    assert verification["method"] == "browser-confirmed"
    assert verification["claim"].startswith("参数 name")
    assert verification["expected"] and verification["actual"]

    template = tmp_path / "default.docx"
    make_default_template.main(["--out", str(template)])
    out = render_docx(context.as_template_context(), template, tmp_path / "out.docx")
    from docx import Document

    texts = [p.text for p in Document(str(out)).paragraphs]
    assert any("验证声明（claim）：参数 name" in t for t in texts)
    assert any("实际结果（actual）：" in t for t in texts)


def test_legacy_sqli_verification_renders_without_four_part(
    report_evidence_dir, tmp_path
):
    """向后兼容：sqli 旧记录无四段式字段 → 默认模板渲染不炸、四段式段落不出现。"""
    context = build_context(report_evidence_dir)
    confirmed = next(f for f in context.confirmed_findings if f.id == "F-2026-0001")
    assert confirmed.verification["method"] == "sqlmap-confirmed"
    assert confirmed.verification["claim"] is None  # 键恒在、值为 null

    template = tmp_path / "default.docx"
    make_default_template.main(["--out", str(template)])
    out = render_docx(context.as_template_context(), template, tmp_path / "out.docx")
    from docx import Document

    texts = [p.text for p in Document(str(out)).paragraphs]
    assert any("验证方法" in t or "sqlmap-confirmed" in t for t in _table_texts(out))
    assert not any("验证声明（claim）" in t for t in texts)


def _table_texts(path):
    from docx import Document

    return [c.text for t in Document(str(path)).tables for r in t.rows for c in r.cells]


# ---- 多 verify skill 接线（runner 层，TestClient 半集成） ----


class FakePhasesVerifyMulti:
    """暴露 verify_skills 的假阶段执行器（M8b 多 verify skill 接口）。"""

    scan_skills = [("web-scan", "L1")]
    verify_skill = "verify-sqli"
    verify_risk_level = "L2"
    verify_skills = [("verify-sqli", "L2"), ("verify-xss", "L2")]

    def __init__(self, runtime):
        self.dir = runtime.engagement.dir
        self.target = runtime.engagement.target
        self.audit = runtime.audit
        self.verified: list[str] = []

    def scan_with_skill(self, skill_name, targets):
        pass

    def triage(self):
        store = FindingStore(self.dir / "findings.jsonl")
        created = []
        for vuln_type, asset, param in (
            ("sqli", f"{self.target}/vulnerabilities/sqli/?id=1", "id"),
            ("xss", f"{self.target}/vulnerabilities/xss_r/?name=1", "name"),
        ):
            finding = Finding(
                id=store.next_id(),
                state=FindingState.SIGNAL,
                vuln_type=vuln_type,
                severity="medium",
                asset=asset,
                param=param,
                confidence="low",
                evidence_kinds=["crawl-endpoint"],
                dedup_key=compute_dedup_key(asset, vuln_type, param),
                created_at="2026-08-14T00:00:00.000+00:00",
                updated_at="2026-08-14T00:00:00.000+00:00",
                audit=self.audit,
            )
            finding.transition(
                FindingState.HYPOTHESIS, actor="triage", reason="fake triage"
            )
            store.append(finding)
            created.append(finding)
        return created

    def verify_covered_vuln_types(self, skill_name=None):
        return {
            "verify-sqli": frozenset({"sqli"}),
            "verify-xss": frozenset({"xss"}),
        }[skill_name or "verify-sqli"]

    def verify_with_skill(self, skill_name):
        self.verified.append(skill_name)
        return []

    def verify(self):
        raise AssertionError("多 skill 接口下不得走旧式 verify()")


@pytest.fixture
def api_workspace(tmp_path):
    (tmp_path / "scope.yaml").write_text("networks: [127.0.0.0/8]\n", encoding="utf-8")
    templates_dir = tmp_path / "templates"
    templates_dir.mkdir()
    shutil.copyfile(
        REPO_ROOT / "templates" / "default_template.docx",
        templates_dir / "default_template.docx",
    )
    return tmp_path


def _wait_pending_conf(client, eng_id, timeout=15.0):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        state = client.get(f"/api/engagements/{eng_id}").json()["state"]
        if state in ("done", "failed"):
            return None
        confs = client.get(f"/api/engagements/{eng_id}/confirmations").json()[
            "confirmations"
        ]
        if confs:
            return confs[0]
        time.sleep(0.05)
    raise AssertionError("等待待确认动作超时")


def test_runner_multi_verify_skills_gating(api_workspace):
    """runner 逐 verify skill 过闸：sqli 归 verify-sqli、xss 归 verify-xss，
    两 phase 均按 skill 执行（确认/审计粒度到 skill）。"""
    instances = []

    def _factory(runtime):
        phases = FakePhasesVerifyMulti(runtime)
        instances.append(phases)
        return phases

    app = create_app(api_workspace, phases_factory=_factory, confirm_timeout=30.0)
    with TestClient(app) as client:
        resp = client.post(
            "/api/engagements",
            json={
                "target": "http://127.0.0.1:8080",
                "scope_paths": ["scope.yaml"],
                "autonomy_mode": "semi_auto",
            },
        )
        assert resp.status_code == 201, resp.text
        eng_id = resp.json()["id"]
        assert client.post(f"/api/engagements/{eng_id}/run").status_code == 202

        seen_actions = []
        while True:
            conf = _wait_pending_conf(client, eng_id)
            if conf is None:
                break
            seen_actions.append((conf["action"], conf["summary"]))
            assert client.post(
                f"/api/confirmations/{conf['cid']}/approve",
                json={"operator": "tester"},
            ).status_code == 200

        deadline = time.monotonic() + 15.0
        while time.monotonic() < deadline:
            if client.get(f"/api/engagements/{eng_id}").json()["state"] == "done":
                break
            time.sleep(0.05)

        phases = instances[0]
        assert phases.verified == ["verify-sqli", "verify-xss"]
        actions = [action for action, _ in seen_actions]
        assert actions == ["verify-sqli", "verify-xss"]  # 逐 skill 各一次确认
        assert any("xss" in summary for _, summary in seen_actions)
